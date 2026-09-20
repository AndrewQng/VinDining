import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def parse_markdown_text(paragraph, text):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.font.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(194, 65, 12)
        else:
            paragraph.add_run(token)

def render_table(doc, table_lines):
    rows_data = []
    for tline in table_lines:
        cells = [c.strip() for c in tline.split('|')[1:-1]]
        if all(re.match(r'^:?-+:?$', c) for c in cells):
            continue
        rows_data.append(cells)
    if not rows_data:
        return
    num_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows_data):
        for c_idx in range(num_cols):
            cell = table.cell(r_idx, c_idx)
            val = row[c_idx] if c_idx < len(row) else ''
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            parse_markdown_text(p, val)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx == 0:
                set_cell_background(cell, '1E3A8A')
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(10)
            else:
                set_cell_background(cell, 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF')
                for run in p.runs:
                    run.font.size = Pt(9.5)
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)

def export():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    md_path = os.path.join(script_dir, "02_BAO_CAO_KHAO_SAT.md")
    docx_path = os.path.join(script_dir, "02_BAO_CAO_KHAO_SAT.docx")
    base_dir = script_dir

    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(51, 51, 51)

    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.splitlines()
    in_code_block = False
    table_lines = []
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith('```'):
            in_code_block = not in_code_block
            i += 1
            continue
        if in_code_block:
            i += 1
            continue
        if line.startswith('|') and line.endswith('|'):
            table_lines.append(line)
            i += 1
            continue
        elif table_lines:
            render_table(doc, table_lines)
            table_lines = []
        if not line:
            i += 1
            continue
        if line == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="E0E0E0"/></w:pBdr>')
            p._element.get_or_add_pPr().append(pBdr)
            i += 1
            continue
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if img_match:
            caption, rel_path = img_match.groups()
            abs_img_path = os.path.normpath(os.path.join(base_dir, rel_path))
            if os.path.exists(abs_img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(12)
                p_img.paragraph_format.space_after = Pt(4)
                run_img = p_img.add_run()
                run_img.add_picture(abs_img_path, width=Inches(6.0))
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(2)
                p_cap.paragraph_format.space_after = Pt(10)
                run_cap = p_cap.add_run(f'Hình: {caption}')
                run_cap.font.italic = True
                run_cap.font.size = Pt(9.5)
                run_cap.font.color.rgb = RGBColor(120, 120, 120)
            i += 1
            continue
        if line.startswith('> 📐'):
            p_note = doc.add_paragraph()
            p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_note.paragraph_format.space_after = Pt(8)
            clean_note = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', line).replace('> 📐', '📐').strip()
            run = p_note.add_run(clean_note)
            run.font.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(100, 100, 150)
            i += 1
            continue
        if line.startswith('# '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(line[2:].strip())
            run.font.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue
        if line.startswith('## '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(line[3:].strip())
            run.font.bold = True
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(37, 99, 235)
            i += 1
            continue
        if line.startswith('### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(line[4:].strip())
            run.font.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(55, 65, 81)
            i += 1
            continue
        if line.startswith('#### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            run = h.add_run(line[5:].strip())
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(75, 85, 99)
            i += 1
            continue
        if line.startswith('* ') or line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            parse_markdown_text(p, line[2:].strip())
            i += 1
            continue
        if line.startswith('  - ') or line.startswith('  * '):
            p = doc.add_paragraph(style='List Bullet 2')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            parse_markdown_text(p, line[4:].strip())
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        parse_markdown_text(p, line)
        i += 1

    if table_lines:
        render_table(doc, table_lines)

    try:
        doc.save(docx_path)
        print("SUCCESS: Exported 02_BAO_CAO_KHAO_SAT.docx successfully.")
    except PermissionError:
        fallback_path = os.path.join(script_dir, "02_BAO_CAO_KHAO_SAT_UPDATED.docx")
        doc.save(fallback_path)
        print(f"NOTICE: 02_BAO_CAO_KHAO_SAT.docx is currently open in Word. Saved to 02_BAO_CAO_KHAO_SAT_UPDATED.docx instead.")

if __name__ == "__main__":
    export()
